"""LexiSure — legal document and compliance education prototype.
Run: streamlit run app.py
"""
import io
import os
from pathlib import Path

import streamlit as st
from dotenv import load_dotenv
from groq import APIError, Groq
from PIL import Image, UnidentifiedImageError
import fitz
import pytesseract

try:
    from docx import Document
except ImportError:
    Document = None

load_dotenv(Path(__file__).with_name(".env"), override=False)
st.set_page_config(page_title="LexiSure | Legal Clarity Studio", page_icon="⚖️", layout="wide")


def setting(name, default=""):
    value = os.getenv(name, "").strip()
    if value:
        return value
    try:
        return str(st.secrets.get(name, default)).strip()
    except Exception:
        return default


KEY = setting("GROQ_API_KEY")
MODEL = setting("GROQ_MODEL", "openai/gpt-oss-20b")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@400;500;600;700;800&display=swap');
:root{color-scheme:dark}.stApp{background:radial-gradient(ellipse at 84% 0%,#19352f 0%,#101920 36%,#090e14 78%);color:#e9f1f2}
[data-testid="stHeader"]{background:transparent}[data-testid="stSidebar"]{background:#101820;border-right:1px solid #29343a}
.block-container{max-width:1320px;padding:2.2rem 2.3rem 4rem}h1,h2,h3{font-family:'Manrope',sans-serif;letter-spacing:-.6px;color:#f0f5f2!important}
p,li,label{font-family:'DM Sans',sans-serif;line-height:1.65}.hero{padding:30px 34px;border:1px solid #3c4b3c;border-radius:24px;background:linear-gradient(115deg,#202d26,#121e25 72%);margin-bottom:22px}
.kicker{color:#b9d79a;font-size:11px;letter-spacing:2px;font-weight:700;text-transform:uppercase}.hero h1{font-size:40px;margin:10px 0}.hero p{color:#b6c3c0;max-width:800px;margin:0}
.pill{display:inline-block;border:1px solid #52634b;color:#d3e6be;background:#273425;border-radius:22px;padding:5px 10px;font-size:11px;margin:14px 6px 0 0}.card{background:#141e25;border:1px solid #2b3a40;border-radius:16px;padding:18px;margin-bottom:12px}.muted{color:#a9b7b9;font-size:13px}
div.stButton>button{background:#25372d;color:#e7f1df;border:1px solid #506446;border-radius:11px;font-weight:650}div.stButton>button:hover{border-color:#bad59a;color:white}
[data-testid="stFileUploader"]{background:#131e24;border:1px dashed #627052;border-radius:14px}[data-testid="stChatMessage"]{background:#141f25;border:1px solid #2b3b41;border-radius:15px}hr{border-color:#2c393e}
</style>
""", unsafe_allow_html=True)

if "legal_messages" not in st.session_state:
    st.session_state.legal_messages = []

with st.sidebar:
    st.markdown("## ⚖️ LexiSure")
    st.caption("LEGAL CLARITY STUDIO · HACKATHON PROTOTYPE")
    st.divider()
    jurisdiction = st.text_input("Jurisdiction", placeholder="e.g., India / state or country")
    language = st.selectbox("Reply language", ["English", "Hinglish", "Hindi", "Telugu"], index=1)
    st.warning("Educational information only—not a lawyer or legal advice. Confirm important decisions with a qualified lawyer in your jurisdiction.")
    st.caption("Files and questions are sent to Groq for requested AI processing. Avoid uploading confidential, privileged, or identifying information.")
    if st.button("Clear conversation", use_container_width=True):
        st.session_state.legal_messages = []
        st.rerun()

st.markdown("""
<div class="hero">
 <div class="kicker">Understand before you sign</div>
 <h1>Legal language, made clearer.</h1>
 <p>Review contract wording, ask compliance questions, and create a first-draft template—with visible limits and risks to verify.</p>
 <span class="pill">Document review</span><span class="pill">Risk scan</span><span class="pill">Compliance Q&A</span><span class="pill">Draft generator</span>
</div>
""", unsafe_allow_html=True)

if not KEY:
    st.info("Add GROQ_API_KEY to Streamlit Secrets to enable AI features.")
client = Groq(api_key=KEY, timeout=60, max_retries=0) if KEY else None


def ask_ai(system, prompt, history=None, max_tokens=1800):
    if not client:
        raise RuntimeError("GROQ_API_KEY is missing. Add it in Streamlit Secrets.")
    messages = [{"role": "system", "content": system}]
    if history:
        messages.extend(history[-10:])
    messages.append({"role": "user", "content": prompt})
    result = client.chat.completions.create(model=MODEL, messages=messages, temperature=0.2, max_completion_tokens=max_tokens)
    answer = result.choices[0].message.content
    if not answer or not answer.strip():
        raise RuntimeError("The AI returned no answer. Try again.")
    return answer.strip()


def error_notice(error):
    status = getattr(error, "status_code", None)
    if status == 401:
        st.error("Groq rejected the key. Check Streamlit Secrets.")
    elif status == 429:
        st.error("Groq request limit reached. Wait a little and retry.")
    elif status in (400, 403, 404):
        st.error("Model unavailable. Check GROQ_MODEL and account access.")
    else:
        st.error("AI request failed. Check deployment logs and try again.")


def extract_document(upload):
    raw = upload.getvalue()
    suffix = Path(upload.name).suffix.lower()
    if len(raw) > 12 * 1024 * 1024:
        raise ValueError("File exceeds 12 MB. Upload a smaller copy.")
    if suffix == ".txt":
        return raw.decode("utf-8", errors="replace"), "Text file"
    if suffix == ".docx":
        if Document is None:
            raise ValueError("DOCX reader is not installed. Check requirements.txt.")
        doc = Document(io.BytesIO(raw))
        return "\n".join(p.text for p in doc.paragraphs), "DOCX text"
    if suffix == ".pdf":
        pdf = fitz.open(stream=raw, filetype="pdf")
        parts = [page.get_text("text") for page in pdf]
        text = "\n".join(parts).strip()
        if len(text) >= 80:
            return text, "PDF text extraction"
        ocr_pages = []
        for page in list(pdf)[:6]:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
            ocr_pages.append(pytesseract.image_to_string(image))
        return "\n".join(ocr_pages), "Scanned PDF OCR (first 6 pages max)"
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        try:
            image = Image.open(io.BytesIO(raw)).convert("RGB")
        except UnidentifiedImageError as exc:
            raise ValueError("This image format could not be read.") from exc
        return pytesseract.image_to_string(image), "Image OCR"
    raise ValueError("Use a PDF, DOCX, TXT, PNG, JPG, or TIFF file.")


base_safety = f"""You are LexiSure, an educational legal-information assistant. Reply in {language}. The stated jurisdiction is {jurisdiction or 'not specified'}; ask for it when it materially changes the answer. You are not a lawyer and must not claim attorney-client privilege, legal representation, or a definitive legal conclusion. Laws can change: distinguish general information from jurisdiction-specific claims, state uncertainty, and tell the user to verify current law with an official source or qualified lawyer. Never fabricate statutes, regulations, case names, citations, dates, or legal requirements. If reliable authority is not in the prompt, say you cannot verify it. Do not advise evasion, fraud, or concealment. Keep responses clear, practical, and appropriately cautious. This prototype has no connected case-law database or live legal-source retrieval."""

tab_review, tab_chat, tab_draft, tab_voice = st.tabs(["📑 Contract Review", "💬 Legal Q&A", "📝 Draft Generator", "🎙️ Voice Notes"])

with tab_review:
    st.subheader("Review a contract or compliance document")
    st.caption("Upload a document to extract its text, explain key clauses, and flag issues to discuss with counsel.")
    target = st.text_input("What should the review focus on?", placeholder="e.g., termination, payment, data processing, IP ownership")
    upload = st.file_uploader("Choose a document", type=["pdf", "docx", "txt", "png", "jpg", "jpeg", "tif", "tiff", "bmp"], key="legal_upload")
    if upload and st.button("Analyze document", type="primary", disabled=not bool(KEY)):
        try:
            with st.spinner("Extracting and reviewing the document…"):
                doc_text, extraction_method = extract_document(upload)
                if len(doc_text.strip()) < 80:
                    raise ValueError("I could not extract enough text. Try a clearer scan or text-based document.")
                st.caption(f"Text extraction: {extraction_method}. Up to 18,000 characters are sent for analysis.")
                report = ask_ai(
                    base_safety + " For contract review, use these headings: Plain-language summary; Key obligations (who must do what and when); Risk level (preliminary 0–100 with rationale, not a legal score); Clauses needing attention; Possible missing topics to ask counsel about; Questions before signing. Quote short exact phrases from the document for every flagged issue; do not invent missing clauses as facts. Distinguish text observed from inference. State that enforceability depends on law and facts.",
                    f"JURISDICTION: {jurisdiction or 'not specified'}\nREVIEW FOCUS: {target or 'general review'}\nDOCUMENT TEXT:\n{doc_text[:18000]}",
                    max_tokens=2200,
                )
            st.markdown(report)
        except APIError as e:
            error_notice(e)
        except Exception as e:
            st.error(str(e))

with tab_chat:
    st.subheader("Ask a legal or compliance question")
    st.caption("Ask about contract language, business setup, employment policies, privacy, intellectual property, or compliance concepts.")
    for msg in st.session_state.legal_messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])
    question = st.chat_input("What would you like explained?", disabled=not bool(KEY), max_chars=3000)
    if question:
        history = st.session_state.legal_messages.copy()
        st.session_state.legal_messages.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)
        try:
            with st.chat_message("assistant"):
                with st.spinner("Preparing a careful explanation…"):
                    answer = ask_ai(base_safety, question, history, max_tokens=1200)
                st.markdown(answer)
            st.session_state.legal_messages.append({"role": "assistant", "content": answer})
        except APIError as e:
            st.session_state.legal_messages.pop()
            error_notice(e)
        except Exception as e:
            st.session_state.legal_messages.pop()
            st.error(str(e))

with tab_draft:
    st.subheader("Create a first-draft document")
    st.caption("Generate a discussion draft only. Have a qualified lawyer adapt and review it before use.")
    kind = st.selectbox("Document type", ["Non-disclosure agreement (NDA)", "Memorandum of understanding (MOU)", "Freelance services agreement", "Partnership discussion draft", "Privacy policy outline", "Other"])
    party_a = st.text_input("Party A / organization", placeholder="Use placeholders if preferred")
    party_b = st.text_input("Party B", placeholder="Use placeholders if preferred")
    terms = st.text_area("Key terms and context", height=150, placeholder="Scope, duration, payment, confidentiality, termination, governing law, and anything else to include")
    if st.button("Generate discussion draft", type="primary", disabled=not bool(KEY)):
        if not terms.strip():
            st.info("Add a short description of the intended terms first.")
        else:
            try:
                with st.spinner("Preparing a clearly labeled draft…"):
                    draft = ask_ai(
                        base_safety + " Draft a plainly labeled, non-final discussion template. Use bracketed placeholders for unknown facts. Never invent jurisdiction requirements, legal citations, or assert enforceability. Include a prominent header 'DRAFT FOR LAWYER REVIEW — NOT READY TO SIGN', numbered sections, and a checklist of points requiring local counsel review. Do not create deceptive or unlawful provisions.",
                        f"Document type: {kind}\nJurisdiction: {jurisdiction or 'unspecified'}\nParty A: {party_a or '[Party A]'}\nParty B: {party_b or '[Party B]'}\nRequested terms:\n{terms}",
                        max_tokens=2200,
                    )
                st.markdown(draft)
                st.download_button("Download draft as text", data=draft, file_name="lexisure-discussion-draft.txt", mime="text/plain")
            except APIError as e:
                error_notice(e)
            except Exception as e:
                st.error(str(e))

with tab_voice:
    st.subheader("Transcribe a voice note")
    st.caption("Record or upload a short note. Transcription is sent to Groq; you can then ask for a plain-language explanation.")
    audio = st.audio_input("Record a legal or compliance question", sample_rate=16000)
    if audio and st.button("Transcribe voice note", disabled=not bool(KEY)):
        if audio.size > 20 * 1024 * 1024:
            st.error("Audio is over 20 MB. Record a shorter note.")
        else:
            try:
                with st.spinner("Transcribing…"):
                    transcript = client.audio.transcriptions.create(
                        file=(audio.name or "voice-note.wav", audio.getvalue()),
                        model="whisper-large-v3-turbo",
                        response_format="json",
                        temperature=0.0,
                    )
                st.text_area("Transcript", value=transcript.text or "", height=130)
                st.caption("Transcription may contain errors. Check important names, dates, amounts, and legal terms against the original recording.")
            except APIError as e:
                error_notice(e)
            except Exception as e:
                st.error(f"Transcription failed: {e}")

st.divider()
st.caption("Prototype limitations: no lawyer review, legal-source retrieval, case-law search, trained legal model, or certification for GDPR/HIPAA compliance. Contract analysis can miss issues. Do not upload privileged or highly confidential documents.")
